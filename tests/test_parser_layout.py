# -*- coding: utf-8 -*-
"""파서 점검(2026-09-30)에서 고친 동작. 근거: docs/DESIGN_NOTES.md#파서-점검"""
import io
import zipfile

import fitz
from docx import Document as DocxDocument

from app.plugins.parsers.docx_structured import parse_blocks as docx_blocks
from app.plugins.parsers.hwpx_structured import parse_blocks as hwpx_blocks
from app.plugins.parsers.pymupdf_parser import (PyMuPDFParser, carry_table_header, garbled_text, place_key,
                                                repeated_margin_text)
from app.services.document_processor import sniff_hwp_extension


def _two_column_page():
    """왼쪽 단을 다 쓰고 오른쪽 단을 쓰는, 줄 높이가 서로 맞는 2단 쪽."""
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    for i in range(6):
        page.insert_text((50, 100 + i * 40), f"left column line {i}")
    for i in range(6):
        page.insert_text((320, 100 + i * 40), f"right column line {i}")
    return doc, page


def test_two_column_follows_recorded_order():
    _doc, page = _two_column_page()
    items, _ = PyMuPDFParser._text_items(page)
    text = " ".join(t for _, t, _ in sorted(items, key=lambda e: e[0]))
    assert text.index("left column line 5") < text.index("right column line 0")


def test_figure_goes_after_block_above_in_same_column():
    entries = [(0.0, "left top", [50, 90, 280, 110]), (1.0, "left bottom", [50, 400, 280, 420]),
               (2.0, "right top", [320, 90, 550, 110])]
    key = place_key(entries, [50, 150, 280, 350])
    assert 0.0 < key < 1.0


def test_text_only_touching_a_table_edge_is_kept():
    """soffice PDF 의 굵은 제목 글자 도형이 작은 '표'로 잡히고, 그 표에 걸친 제목이 통째로 사라졌었다."""
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((60, 108), "A long heading whose left end overlaps a small grid of rules")
    for x in (58, 70, 82):
        page.draw_rect(fitz.Rect(x, 99, x + 11, 112), color=(0, 0, 0))
    items, _ = PyMuPDFParser._text_items(page)
    assert any("A long heading whose left end overlaps" in t for _, t, _ in items)


def test_repeated_margin_text_is_found_across_pages():
    doc = fitz.open()
    for n in range(4):
        page = doc.new_page(width=595, height=842)
        page.insert_text((50, 30), "Company internal - confidential")
        page.insert_text((290, 820), f"- {n + 1} -")
        page.insert_text((50, 400), f"body text on page {n + 1}")
    keys = repeated_margin_text(doc)
    items, _ = PyMuPDFParser._text_items(doc[1], keys)
    texts = [t for _, t, _ in items]
    assert texts == ["body text on page 2"]


def test_table_continued_on_next_page_gets_header_back():
    prev = "|번호|시설명|\n|---|---|\n|35|공연장|"
    cur = [(3.0, "|36|체육관|\n|---|---|\n|37|수영장|", [50, 60, 500, 200])]
    entries, last = carry_table_header(cur, prev)
    assert entries[0][1].splitlines() == ["|번호|시설명|", "|---|---|", "|36|체육관|", "|37|수영장|"]
    assert last == entries[0][1]
    same, _ = carry_table_header([(0.0, "|번호|시설명|\n|---|---|\n|36|체육관|", [0, 0, 1, 1])], prev)
    assert same[0][1].count("번호") == 1  # 머리행을 이미 반복한 표는 그대로


def test_garbled_text_layer():
    assert garbled_text("ײ˯ࡿࡆԫ z лː୚ࢄ୬ɹ࢝ਹٷْ ʵ५Æ ۺ࢝ࢲଙ z ۗԸࡈی߳ିֵ˒ л ࣵ˯ ̘ܺʾࢧ ײ˯ࡿ࢒ݣʀ ߻μऌ৊ֱ z ܎ࢇࠚࢇࡿ۟ی˒")
    assert not garbled_text("한빛시 시설관리공단은 올해 상반기 동안 공공체육시설 열두 곳의 운영 실태를 점검했다. PDF 2025")
    assert not garbled_text("مجلة البحث العلمي في التربية العدد الحادي والعشرون الجزء الثاني عشر")
    assert not garbled_text("東京都の人口は約千四百万人です。カタカナとひらがなが混ざった文章です。ABC")


def _docx_table_markdown(build):
    d = DocxDocument()
    build(d)
    buf = io.BytesIO()
    d.save(buf)
    with zipfile.ZipFile(buf) as z:
        blocks = docx_blocks(z.read("word/document.xml"), {})
    return next(b["markdown"] for b in blocks if b["type"] == "table")


def test_docx_merged_cells_keep_columns_and_fill_values():
    def build(d):
        t = d.add_table(rows=3, cols=3)
        t.cell(0, 0).merge(t.cell(0, 1)).text = "group"
        t.cell(0, 2).text = "note"
        t.cell(1, 0).merge(t.cell(2, 0)).text = "dept"
        t.cell(1, 1).text = "a"
        t.cell(1, 2).text = "1"
        t.cell(2, 1).text = "b"
        t.cell(2, 2).text = "2"

    rows = [r for r in _docx_table_markdown(build).splitlines() if "---" not in r]
    assert rows[0] == "| group | group | note |"
    assert rows[2] == "| dept | b | 2 |"


HWPX_NS = ('xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section" '
           'xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph"')


def _tc(r, c, text, rs=1, cs=1):
    return (f'<hp:tc><hp:subList><hp:p><hp:run><hp:t>{text}</hp:t></hp:run></hp:p></hp:subList>'
            f'<hp:cellAddr colAddr="{c}" rowAddr="{r}"/><hp:cellSpan colSpan="{cs}" rowSpan="{rs}"/></hp:tc>')


def test_hwpx_spans_filled_and_header_not_in_body(tmp_path):
    section = (
        f'<hs:sec {HWPX_NS}>'
        '<hp:p><hp:run><hp:ctrl><hp:header><hp:subList><hp:p><hp:run><hp:t>RUNNING HEADER</hp:t></hp:run></hp:p>'
        '</hp:subList></hp:header></hp:ctrl><hp:t>Title</hp:t></hp:run></hp:p>'
        '<hp:p><hp:run><hp:tbl>'
        f'<hp:tr>{_tc(0, 0, "dept", rs=2)}{_tc(0, 1, "a")}</hp:tr>'
        f'<hp:tr>{_tc(1, 1, "b")}</hp:tr>'
        '</hp:tbl></hp:run></hp:p></hs:sec>')
    path = tmp_path / "t.hwpx"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("Contents/section0.xml", section)
    blocks = hwpx_blocks(str(path), describe=False)
    assert blocks[0]["text"] == "Title"
    table = next(b["text"] for b in blocks if b["kind"] == "table")
    assert "| dept | b |" in table


def test_hwp_binary_with_hwpx_extension_is_sniffed():
    assert sniff_hwp_extension(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1rest", ".hwpx") == ".hwp"
    assert sniff_hwp_extension(b"PK\x03\x04rest", ".hwp") == ".hwpx"
    assert sniff_hwp_extension(b"%PDF", ".pdf") == ".pdf"
