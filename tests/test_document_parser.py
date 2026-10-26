# -*- coding: utf-8 -*-
"""process-gpt-document-parser 출력이 memento 형식(PDF 쪽별 blocks_json, 흐르는 문서 블록)으로 바뀐다."""
import asyncio
import io
import json

import docx
import pytest

from app.plugins.parsers import get_pdf_parser, vision
from app.services.document_blocks import BLOCKS_KEY, build_blocks
from app.services.document_processor import DocumentProcessor


@pytest.fixture(autouse=True)
def no_vlm(monkeypatch):
    monkeypatch.setattr(vision, "PDF_VISION_ENABLED", False)


def _pdf(lines):
    """Helvetica 글 줄만 있는 한 쪽짜리 PDF."""
    ops = "".join(f"BT /F1 12 Tf 72 {720 - i * 40} Td ({t}) Tj ET\n" for i, t in enumerate(lines)).encode()
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(ops) + ops + b"endstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, 1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % i + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1))
    out.write(b"".join(b"%010d 00000 n \n" % o for o in offsets))
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%EOF\n" % (len(objs) + 1, xref))
    return out.getvalue()


def test_pdf_page_blocks_point_at_their_text():
    pages = asyncio.run(get_pdf_parser().parse(_pdf(["First paragraph here", "Second paragraph far below"]), "a.pdf"))

    page = pages[0]
    assert page.page_content.startswith("# 페이지 1\n\n")
    spots = json.loads(page.metadata["blocks_json"])
    texts = [page.page_content[s["offset"]:s["offset"] + s["length"]] for s in spots]
    assert "First paragraph here" in " ".join(texts)
    assert all(s["bbox"] for s in spots)
    assert [b["page_number"] for b in build_blocks(pages)] == [1] * len(build_blocks(pages))


def test_docx_headings_and_tables_become_blocks():
    d = docx.Document()
    d.add_heading("제1장 총칙", level=1)
    d.add_paragraph("본문 문단")
    table = d.add_table(rows=2, cols=2)
    for r, row in enumerate([["구분", "값"], ["가", "1"]]):
        for c, text in enumerate(row):
            table.cell(r, c).text = text
    buf = io.BytesIO()
    d.save(buf)
    buf.seek(0)

    docs = asyncio.run(DocumentProcessor().load_document(buf, "a.docx"))

    blocks = docs[0].metadata[BLOCKS_KEY]
    assert blocks[0] == {"kind": "paragraph", "text": "제1장 총칙", "heading_level": 1}
    assert blocks[1]["text"] == "본문 문단"
    assert blocks[2]["kind"] == "table" and "| 가 | 1 |" in blocks[2]["text"]
    assert docs[0].metadata["file_type"] == "docx"


def test_hwp_extension_on_docx_content_is_read_by_content():
    d = docx.Document()
    d.add_paragraph("확장자가 틀린 문서")
    buf = io.BytesIO()
    d.save(buf)
    buf.seek(0)

    docs = asyncio.run(DocumentProcessor().load_document(buf, "a.hwp"))

    assert "확장자가 틀린 문서" in docs[0].page_content
    assert docs[0].metadata["file_type"] == "docx"
